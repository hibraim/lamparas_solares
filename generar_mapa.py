import os
import csv
import re
from docx import Document
from PIL import Image
from io import BytesIO
import folium
from folium.plugins import MarkerCluster

def limpiar_texto(texto):
    if not texto:
        return ""
    return texto.strip().replace('\n', ' ')

def extraer_coordenadas(texto_geo):
    """Extrae latitud y longitud de un texto usando expresiones regulares."""
    if not texto_geo:
        return None, None
    match = re.search(r'(-?\d+\.\d+)\s*,\s*(-?\d+\.\d+)', texto_geo)
    if match:
        return float(match.group(1)), float(match.group(2))
    return None, None

def procesar_word_y_generar_mapa(docx_path="datos.docx", output_dir="salida", img_dir="imagenes"):
    os.makedirs(output_dir, exist_ok=True)
    os.makedirs(img_dir, exist_ok=True)

    if not os.path.exists(docx_path):
        print(f"❌ Error: No se encontró el archivo '{docx_path}' en la ruta especificada.")
        return

    print("📄 Analizando el documento Word...")
    doc = Document(docx_path)

    ubicaciones = []
    errores = []

    # Recolectar imágenes inline del documento
    imagenes_extraidas = []
    for rel_id, rel in doc.part.rels.items():
        if "image" in rel.target_ref:
            try:
                img_blob = rel.target_part.blob
                imagenes_extraidas.append((rel_id, img_blob))
            except Exception as e:
                print(f"⚠️ Advertencia: No se pudo extraer una imagen interna: {e}")

    # Extraer filas de las tablas
    filas_datos = []
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            texto_fila = [cell.text.strip() for cell in row.cells]
            if i == 0 and any("LUGAR" in t.upper() or "GEOLOCALIZACION" in t.upper() for t in texto_fila):
                continue
            if len(row.cells) >= 2:
                lugar = limpiar_texto(row.cells[0].text)
                geo_text = limpiar_texto(row.cells[1].text)
                
                if lugar or geo_text:
                    filas_datos.append((lugar, geo_text))

    # Fallback si no hay tablas estructuradas
    if not filas_datos:
        for p in doc.paragraphs:
            lat, lon = extraer_coordenadas(p.text)
            if lat and lon:
                filas_datos.append(("Ubicación extraída", p.text))

    print(f"🔍 Se encontraron {len(filas_datos)} registros y {len(imagenes_extraidas)} imágenes.")

    for idx, (lugar, geo_text) in enumerate(filas_datos):
        lat, lon = extraer_coordenadas(geo_text)
        
        if lat is None or lon is None:
            errores.append({
                "indice": idx + 1,
                "lugar": lugar if lugar else "Desconocido",
                "datos_originales": geo_text,
                "motivo": "Coordenadas inválidas o faltantes"
            })
            continue

        img_filename = None
        if idx < len(imagenes_extraidas):
            _, img_blob = imagenes_extraidas[idx]
            try:
                img = Image.open(BytesIO(img_blob))
                if img.mode in ("RGBA", "P"):
                    img = img.convert("RGB")
                
                img_filename = f"ubicacion_{idx + 1}.jpg"
                img_path = os.path.join(img_dir, img_filename)
                
                # Compresión optimizada para mantener bajo peso (<10 MB total)
                img.thumbnail((500, 500))
                img.save(img_path, "JPEG", quality=65)
            except Exception as e:
                print(f"⚠️ Error procesando imagen para {lugar}: {e}")

        ubicaciones.append({
            "id": idx + 1,
            "nombre": lugar if lugar else f"Luminaria #{idx + 1}",
            "lat": lat,
            "lon": lon,
            "imagen": f"../{img_dir}/{img_filename}" if img_filename else None
        })

    if errores:
        with open("errores.csv", mode="w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["indice", "lugar", "datos_originales", "motivo"])
            writer.writeheader()
            writer.writerows(errores)
        print(f"⚠️ Se detectaron {len(errores)} registros incompletos. Guardados en 'errores.csv'.")

    if not ubicaciones:
        print("❌ Error crítico: No hay ubicaciones válidas para procesar.")
        return

    print("🗺️ Generando mapa interactivo con lista unificada y buscador...")
    
    centro_lat = sum(u["lat"] for u in ubicaciones) / len(ubicaciones)
    centro_lon = sum(u["lon"] for u in ubicaciones) / len(ubicaciones)

    # Crear mapa base con Esri
    m = folium.Map(
        location=[centro_lat, centro_lon],
        zoom_start=11,
        tiles="Esri.WorldStreetMap"
    )

    markers_js = ""
    sidebar_items_html = ""

    for u in ubicaciones:
        img_tag = f"""
            <div style="text-align: center; margin-bottom: 8px;">
                <img src="{u['imagen']}" onclick="window.open('{u['imagen']}', '_blank')" 
                     style="width: 100%; max-height: 140px; object-fit: cover; border-radius: 6px; cursor: pointer;" 
                     title="Haz clic para ampliar" alt="Evidencia">
                <span style="font-size: 9px; color: #64748b; display: block; margin-top: 2px;">🔍 Clic para ampliar</span>
            </div>
        """ if u["imagen"] else """
            <div style="background: #f1f3f5; padding: 10px; text-align: center; border-radius: 6px; color: #6c757d; font-size: 11px; margin-bottom: 8px;">
                📷 Sin fotografía
            </div>
        """

        google_maps_url = f"https://www.google.com/maps/dir/?api=1&destination={u['lat']},{u['lon']}"
        
        popup_html = f"""
        <div style="font-family: 'Segoe UI', sans-serif; width: 220px;">
            {img_tag}
            <h4 style="margin: 0 0 4px 0; color: #0f172a; font-size: 13px; font-weight: 600;">{u['nombre']}</h4>
            <p style="margin: 0 0 6px 0; color: #475569; font-size: 11px;"><b>Coord:</b> {u['lat']:.4f}, {u['lon']:.4f}</p>
            <a href="{google_maps_url}" target="_blank" style="display: block; width: 100%; background-color: #16a34a; color: white; text-align: center; padding: 5px 0; border-radius: 4px; text-decoration: none; font-weight: bold; font-size: 10px;">
                📍 CÓMO LLEGAR
            </a>
        </div>
        """

        markers_js += f"""
        var marker_{u['id']} = L.marker([{u['lat']}, {u['lon']}], {{
            icon: L.divIcon({{
                className: 'custom-pin',
                html: '<div style="background-color: #2563eb; width: 24px; height: 24px; border-radius: 50%; border: 2px solid white; display: flex; align-items: center; justify-content: center; color: white; font-size: 10px; box-shadow: 0 2px 5px rgba(0,0,0,0.3);"><i class="fa fa-bolt"></i></div>',
                iconSize: [24, 24],
                iconAnchor: [12, 12]
            }})
        }}).addTo(mainCluster);
        
        marker_{u['id']}.bindPopup({repr(popup_html)});
        markersMap[{u['id']}] = marker_{u['id']};
        """

        sidebar_items_html += f"""
        <div class="ubicacion-item" data-name="{u['nombre'].lower()}" onclick="centrarMapa({u['lat']}, {u['lon']}, {u['id']})">
            <div class="item-name">💡 {u['nombre']}</div>
            <div class="item-sub">Lat: {u['lat']:.4f}, Lon: {u['lon']:.4f}</div>
        </div>
        """

    custom_ui_html = f"""
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/leaflet.markercluster/1.5.1/MarkerCluster.css"/>
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/leaflet.markercluster/1.5.1/MarkerCluster.Default.css"/>
    <script src="https://cdnjs.cloudflare.com/ajax/libs/leaflet.markercluster/1.5.1/leaflet.markercluster.js"></script>

    <style>
        #sidebar {{
            position: absolute;
            top: 0;
            left: 0;
            width: 340px;
            height: 100vh;
            background: #ffffff;
            z-index: 1000;
            box-shadow: 4px 0 15px rgba(0,0,0,0.15);
            display: flex;
            flex-direction: column;
            font-family: 'Segoe UI', Tahoma, sans-serif;
        }}
        .sidebar-header {{
            padding: 16px;
            background: #0f172a;
            color: white;
        }}
        .sidebar-header h2 {{
            margin: 0;
            font-size: 16px;
            font-weight: 700;
        }}
        .sidebar-search {{
            padding: 10px 16px;
            background: #f8fafc;
            border-bottom: 1px solid #e2e8f0;
        }}
        .sidebar-search input {{
            width: 100%;
            padding: 8px 12px;
            border: 1px solid #cbd5e1;
            border-radius: 6px;
            font-size: 13px;
            outline: none;
            box-sizing: border-box;
        }}
        .sidebar-content {{
            overflow-y: auto;
            flex: 1;
            padding: 10px;
        }}
        .ubicacion-item {{
            padding: 10px 12px;
            margin-bottom: 6px;
            border-radius: 6px;
            cursor: pointer;
            transition: background 0.2s;
            border: 1px solid #f1f5f9;
            background: #ffffff;
        }}
        .ubicacion-item:hover {{
            background: #eff6ff;
            border-color: #bfdbfe;
        }}
        .item-name {{
            font-size: 12px;
            font-weight: 600;
            color: #1e293b;
        }}
        .item-sub {{
            font-size: 10px;
            color: #64748b;
            margin-top: 2px;
        }}
        .leaflet-container {{
            width: calc(100% - 340px) !important;
            left: 340px !important;
            position: absolute !important;
            height: 100vh !important;
        }}
        @media (max-width: 768px) {{
            #sidebar {{ width: 100%; height: 40vh; position: relative; }}
            .leaflet-container {{ width: 100% !important; left: 0 !important; height: 60vh !important; }}
        }}
    </style>

    <div id="sidebar">
        <div class="sidebar-header">
            <h2>💡 Alumbrado Público - Lámparas Solares</h2>
            <p style="margin: 4px 0 0 0; font-size: 11px; color: #94a3b8;">Total de elementos: {len(ubicaciones)}</p>
        </div>
        <div class="sidebar-search">
            <input type="text" id="searchInput" placeholder="🔍 Buscar por nombre de lámpara..." onkeyup="filtrarUbicaciones()">
        </div>
        <div class="sidebar-content" id="sidebarList">
            {sidebar_items_html}
        </div>
    </div>

    <script>
        var markersMap = {{}};
        var mapInstance;
        var mainCluster;

        function centrarMapa(lat, lon, id) {{
            if (mapInstance) {{
                mapInstance.setView([lat, lon], 17);
                if (markersMap[id]) {{
                    markersMap[id].openPopup();
                }}
            }}
        }}

        function filtrarUbicaciones() {{
            var input = document.getElementById('searchInput').value.toLowerCase();
            var items = document.getElementsByClassName('ubicacion-item');
            
            for (var i = 0; i < items.length; i++) {{
                var text = items[i].getAttribute('data-name');
                if (text.includes(input)) {{
                    items[i].style.display = "";
                }} else {{
                    items[i].style.display = "none";
                }}
            }}
        }}

        window.addEventListener('load', function() {{
            setTimeout(function() {{
                for (var key in window) {{
                    if (window[key] instanceof L.Map) {{
                        mapInstance = window[key];
                        break;
                    }}
                }}
                
                if (mapInstance) {{
                    mainCluster = L.markerClusterGroup();
                    mapInstance.addLayer(mainCluster);
                    {markers_js}
                }} else {{
                    console.error("No se pudo detectar la instancia del mapa Leaflet.");
                }}
            }}, 300);
        }});
    </script>
    """

    m.get_root().html.add_child(folium.Element(custom_ui_html))

    output_path = os.path.join(output_dir, "mapa.html")
    m.save(output_path)
    print(f"✨ ¡Listo! Mapa generado con las {len(ubicaciones)} lámparas en lista unificada: {output_path}")

if __name__ == "__main__":
    procesar_word_y_generar_mapa()